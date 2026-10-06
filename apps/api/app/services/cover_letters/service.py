import hashlib
import io
import json
import logging
import uuid

import docx
from docx.shared import Pt
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.ai.prompts import cover_letter as prompt
from app.ai.providers import AIProvider
from app.ai.schemas import LetterDraft, LetterVerifyReport
from app.ai.services.runner import LLMRunner, LLMTask
from app.core.config import Settings
from app.models import CandidateProfile, CoverLetter, Job, ResumeVersion
from app.repositories.profiles import ProfileRepository
from app.repositories.settings import SettingsRepository
from app.schemas.cover_letters import (
    CoverLetterCreate,
    LetterContent,
    Paragraph,
    Sentence,
    SentencesUpdate,
)
from app.schemas.profile import ProfileData
from app.schemas.tailoring import ResumeDocument
from app.services.cover_letters.checks import check_claim, check_company, check_connective
from app.services.cover_letters.template import build_template
from app.services.matching.skills import find_skills
from app.services.tailoring.document import base_document, item_texts

logger = logging.getLogger(__name__)

DRAFT = LLMTask(prompt.TASK, prompt.PROMPT_VERSION, prompt.SYSTEM_PROMPT, prompt.MAX_OUTPUT_TOKENS)
VERIFY = LLMTask(
    prompt.VERIFY_TASK,
    prompt.VERIFY_PROMPT_VERSION,
    prompt.VERIFY_SYSTEM_PROMPT,
    prompt.VERIFY_MAX_OUTPUT_TOKENS,
)
UNVERIFIED = "Couldn't be checked by the second validation pass. Try generating again."
NO_AI_NOTICE = "AI processing is off, so this is a template letter built from your resume."


class CoverLetterError(Exception):
    def __init__(self, message: str, status: int = 422) -> None:
        super().__init__(message)
        self.status = status


def check(sentence: Sentence, texts: dict[str, str], job_text: str) -> list[str]:
    if sentence.kind == "claim":
        return check_claim(sentence.text, sentence.sources, texts)
    if sentence.kind == "company":
        return check_company(sentence.text, sentence.job_quote, job_text)
    return check_connective(sentence.text)


class CoverLetterService:
    def __init__(
        self, session: Session, user_id: uuid.UUID, provider: AIProvider | None, settings: Settings
    ) -> None:
        self._session = session
        self._user_id = user_id
        self._provider = provider
        self._settings = settings

    def get(self, letter_id: uuid.UUID) -> CoverLetter | None:
        return self._session.scalar(
            select(CoverLetter).where(
                CoverLetter.id == letter_id, CoverLetter.user_id == self._user_id
            )
        )

    def all(self, job_id: uuid.UUID | None = None) -> list[CoverLetter]:
        stmt = select(CoverLetter).where(CoverLetter.user_id == self._user_id)
        if job_id:
            stmt = stmt.where(CoverLetter.job_id == job_id)
        return list(self._session.scalars(stmt.order_by(CoverLetter.created_at.desc())))

    def _job(self, job_id: uuid.UUID | None) -> Job | None:
        if job_id is None:
            return None
        return self._session.scalar(
            select(Job).where(Job.id == job_id, Job.user_id == self._user_id)
        )

    def _document(self, letter: CoverLetter) -> ResumeDocument:
        """A chosen saved tailored resume version, else the verified profile."""
        if letter.resume_version_id:
            rv = self._session.scalar(
                select(ResumeVersion).where(
                    ResumeVersion.id == letter.resume_version_id,
                    ResumeVersion.user_id == self._user_id,
                )
            )
            if rv and rv.content:
                return ResumeDocument.model_validate(rv.content)
        profile = self._session.scalar(
            select(CandidateProfile).where(
                CandidateProfile.user_id == self._user_id,
                CandidateProfile.version == letter.profile_version,
            )
        )
        return base_document(ProfileData.model_validate(profile.data) if profile else ProfileData())

    @staticmethod
    def paragraphs(letter: CoverLetter) -> list[Paragraph]:
        return [Paragraph.model_validate(p) for p in letter.paragraphs]

    # --- create and generate -------------------------------------------------------------

    def create(self, body: CoverLetterCreate) -> CoverLetter:
        job = self._job(body.job_id)
        if job is None:
            raise CoverLetterError("Job not found", 404)
        profile = ProfileRepository(self._session).current(self._user_id)
        if profile is None or not ProfileData.model_validate(profile.data).experience:
            raise CoverLetterError(
                "Upload your resume first, so the letter has something to draw on."
            )
        if body.resume_version_id is not None:
            rv = self._session.scalar(
                select(ResumeVersion).where(
                    ResumeVersion.id == body.resume_version_id,
                    ResumeVersion.user_id == self._user_id,
                )
            )
            if rv is None or rv.status != "saved":
                raise CoverLetterError("Choose a saved tailored resume version.")
        latest = self._session.scalar(
            select(func.max(CoverLetter.version)).where(CoverLetter.user_id == self._user_id)
        )
        letter = CoverLetter(
            user_id=self._user_id,
            version=(latest or 0) + 1,
            profile_version=profile.version,
            job_id=job.id,
            resume_version_id=body.resume_version_id,
            job_title=job.title,
            company=job.company,
            name=f"Cover letter: {job.title} at {job.company}"[:200],
            tone=body.tone,
            length=body.length,
            status="generating",
            method="template",
            paragraphs=[],
        )
        self._session.add(letter)
        self._session.commit()
        return letter

    def generate(self, letter_id: uuid.UUID) -> None:
        letter = self.get(letter_id)
        if letter is None or letter.status != "generating":
            return
        job = self._job(letter.job_id)
        if job is None:
            letter.status, letter.notice = "failed", "The job was removed."
            self._session.commit()
            return
        doc = self._document(letter)
        job_text = f"{job.title}\n{job.description_text}"
        settings_row = SettingsRepository(self._session).get_for_user(self._user_id)
        consent = bool(settings_row and settings_row.llm_consent)

        paragraphs: list[Paragraph] | None = None
        notice: str | None = NO_AI_NOTICE
        if consent and self._provider is not None:
            paragraphs, notice = self._ai_letter(doc, job, letter)
        if paragraphs is None:
            paragraphs, _ = build_template(
                doc, job.title, job.company, find_skills(job_text), letter.length
            )
            texts = item_texts(doc)
            for p in paragraphs:  # template sentences go through the same code checks
                for s in p.sentences:
                    s.issues = check(s, texts, job_text)
                    if s.issues:
                        s.status, s.included = "unsupported", False
        else:
            letter.method = "ai"
        letter.paragraphs = [p.model_dump(mode="json") for p in paragraphs]
        letter.status, letter.notice = "ready", notice
        self._session.commit()
        sentences = [s for p in paragraphs for s in p.sentences]
        logger.info(
            "cover_letter_generated",
            extra={
                "letter_id": str(letter.id),
                "method": letter.method,
                "sentences": len(sentences),
                "unsupported": sum(s.status == "unsupported" for s in sentences),
            },
        )

    def _ai_letter(
        self, doc: ResumeDocument, job: Job, letter: CoverLetter
    ) -> tuple[list[Paragraph] | None, str | None]:
        runner = LLMRunner(self._session, self._provider, self._settings)
        texts = item_texts(doc)
        items = {
            "name": doc.contact.name,
            "skills": [{"id": s.id, "name": s.name} for s in doc.skills],
            "roles": [
                {
                    "id": r.id,
                    "title": r.title,
                    "company": r.company,
                    "dates": r.date_text,
                    "bullets": [{"id": b.id, "text": b.text} for b in r.bullets],
                }
                for r in doc.experience
            ],
            "projects": [
                {"id": p.id, "name": p.name, "description": p.description} for p in doc.projects
            ],
        }
        posting = {
            "title": job.title,
            "company": job.company,
            "description": job.description_text[:6000],
        }
        key = hashlib.sha256(
            json.dumps(
                [items, job.content_hash, letter.tone, letter.length], sort_keys=True
            ).encode()
        ).hexdigest()
        outcome = runner.run(
            self._user_id,
            DRAFT,
            prompt=prompt.build_prompt(items, letter.tone, letter.length, posting),
            schema=LetterDraft,
            cache_key=key,
            consent=True,
        )
        if outcome.value is None or not outcome.value.paragraphs:
            return (
                None,
                f"{outcome.notice or 'The AI returned no letter.'} "
                "This is a template letter instead.",
            )

        job_text = f"{job.title}\n{job.description_text}"
        paragraphs: list[Paragraph] = []
        for pi, para in enumerate(outcome.value.paragraphs):
            sentences = []
            for si, raw in enumerate(para.sentences):
                s = Sentence(
                    id=f"a{pi}-{si}",
                    text=" ".join(raw.text.split()),
                    kind=raw.kind,
                    sources=[x for x in raw.sources if x in texts],
                    job_quote=raw.job_quote,
                    source="ai",
                    status="proposed",
                )
                s.issues = check(s, texts, job_text)
                if raw.kind == "claim" and len(s.sources) < len(raw.sources):
                    s.issues.append("Cites items that aren't on your resume")
                if s.issues:
                    s.status, s.included = "unsupported", False
                sentences.append(s)
            paragraphs.append(Paragraph(id=f"p{pi}", sentences=sentences))
        self._verify(runner, paragraphs, texts, key)
        return paragraphs, None

    def _verify(
        self, runner: LLMRunner, paragraphs: list[Paragraph], texts: dict[str, str], key: str
    ) -> None:
        """Second pass over every AI sentence that passed the code checks."""
        pending = [s for p in paragraphs for s in p.sentences if s.status == "proposed"]
        if not pending:
            return
        payload = []
        for s in pending:
            evidence = (
                "\n".join(texts[x] for x in s.sources)
                if s.kind == "claim"
                else (s.job_quote or "")
                if s.kind == "company"
                else ""
            )
            payload.append(
                {"sentence_id": s.id, "kind": s.kind, "text": s.text, "evidence": evidence}
            )
        outcome = runner.run(
            self._user_id,
            VERIFY,
            prompt=prompt.build_verify_prompt(payload),
            schema=LetterVerifyReport,
            cache_key=hashlib.sha256(json.dumps([key, payload]).encode()).hexdigest(),
            consent=True,
        )
        results = {r.sentence_id: r for r in outcome.value.results} if outcome.value else {}
        for s in pending:
            r = results.get(s.id)
            if r is None:
                s.status, s.included, s.issues = "unsupported", False, [UNVERIFIED]
            elif not r.supported:
                s.status, s.included = "unsupported", False
                s.issues = [f"Second check: {p}" for p in r.problems] or [
                    "Second check: not supported"
                ]

    # --- edit, preview, save -------------------------------------------------------------

    def update(self, letter: CoverLetter, body: SentencesUpdate) -> CoverLetter:
        if letter.status != "ready":
            raise CoverLetterError("Only a letter under review can be edited.", 409)
        paragraphs = self.paragraphs(letter)
        by_id = {s.id: s for p in paragraphs for s in p.sentences}
        for sentence_id, change in body.updates.items():
            s = by_id.get(sentence_id)
            if s is None:
                raise CoverLetterError(f"Unknown sentence: {sentence_id}")
            if change.text is not None and " ".join(change.text.split()) != s.text:
                # The user's own words: labeled as such, their responsibility, not machine-checked.
                s.text, s.source, s.status, s.issues = (
                    " ".join(change.text.split()),
                    "user",
                    "proposed",
                    [],
                )
                s.included = True
            if change.included is not None:
                if change.included and s.status == "unsupported":
                    raise CoverLetterError(
                        "Unsupported sentences can't be included. "
                        "Edit it into your own words instead."
                    )
                s.included = change.included
        paragraphs_by_id = {p.id: p for p in paragraphs}
        for new in body.add:
            para = paragraphs_by_id.get(new.paragraph_id)
            if para is None:
                raise CoverLetterError(f"Unknown paragraph: {new.paragraph_id}")
            para.sentences.append(
                Sentence(
                    id=f"u{uuid.uuid4().hex[:8]}",
                    text=" ".join(new.text.split()),
                    kind="claim",
                    source="user",
                    status="proposed",
                )
            )
        letter.paragraphs = [p.model_dump(mode="json") for p in paragraphs]
        self._session.commit()
        return letter

    def content(self, letter: CoverLetter) -> LetterContent:
        if letter.status == "saved" and letter.content:
            return LetterContent.model_validate(letter.content)
        paragraphs = [
            Paragraph(
                id=p.id, sentences=[s for s in p.sentences if s.included and s.status == "proposed"]
            )
            for p in self.paragraphs(letter)
        ]
        return LetterContent(
            paragraphs=[p for p in paragraphs if p.sentences],
            signature=self._document(letter).contact.name,
        )

    def preview(self, letter: CoverLetter) -> list[str]:
        return [" ".join(s.text for s in p.sentences) for p in self.content(letter).paragraphs]

    def save(self, letter: CoverLetter, name: str | None) -> CoverLetter:
        if letter.status != "ready":
            raise CoverLetterError("Only a letter under review can be saved.", 409)
        content = self.content(letter)
        if not content.paragraphs:
            raise CoverLetterError("Include at least one sentence before saving.")
        # Defense in depth: re-check every included machine-written sentence.
        job = self._job(letter.job_id)
        texts = item_texts(self._document(letter))
        job_text = f"{job.title}\n{job.description_text}" if job else ""
        for p in content.paragraphs:
            for s in p.sentences:
                if s.source != "user" and check(s, texts, job_text):
                    raise CoverLetterError("An included sentence no longer passes the checks.")
        letter.content = content.model_dump(mode="json")
        letter.status = "saved"
        if name and name.strip():
            letter.name = name.strip()[:200]
        self._session.commit()
        logger.info("cover_letter_saved", extra={"letter_id": str(letter.id)})
        return letter

    def delete(self, letter: CoverLetter) -> None:
        self._session.delete(letter)
        self._session.commit()

    def render_docx(self, letter: CoverLetter) -> bytes:
        content = self.content(letter)
        out = docx.Document()
        out.styles["Normal"].font.name = "Calibri"
        out.styles["Normal"].font.size = Pt(11)
        out.add_paragraph(f"Dear Hiring Team at {letter.company},")
        for p in content.paragraphs:
            out.add_paragraph(" ".join(s.text for s in p.sentences))
        out.add_paragraph("Sincerely,")
        if content.signature:
            out.add_paragraph(content.signature)
        buffer = io.BytesIO()
        out.save(buffer)
        return buffer.getvalue()
