"""Upload validation and text extraction. Resumes are untrusted input."""

import io
import zipfile
from dataclasses import dataclass
from pathlib import PurePath
from typing import Literal

from app.core.config import Settings
from app.utils.text import normalize_text

FileType = Literal["pdf", "docx", "txt"]

_EXTENSIONS: dict[str, FileType] = {".pdf": "pdf", ".docx": "docx", ".txt": "txt"}
MEDIA_TYPES: dict[FileType, str] = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain; charset=utf-8",
}


class UploadRejectedError(Exception):
    """The file is not acceptable. The message is safe to show to the user."""


class ExtractionError(Exception):
    """Text could not be extracted. The message is safe to show to the user."""


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    page_count: int | None


def detect_file_type(data: bytes) -> FileType | None:
    """Identify the type from content, never from the filename."""
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = set(zf.namelist())
        except zipfile.BadZipFile:
            return None
        return "docx" if "word/document.xml" in names and "[Content_Types].xml" in names else None
    if b"\x00" in data[:8192]:
        return None
    try:
        data.decode("utf-8")
    except UnicodeDecodeError:
        return None
    return "txt"


def safe_display_name(filename: str | None) -> str:
    name = PurePath((filename or "resume").replace("\\", "/")).name
    name = "".join(ch for ch in name if ch.isprintable() and ch not in '<>:"/\\|?*').strip()
    return (name or "resume")[:255]


def validate_upload(filename: str | None, data: bytes, settings: Settings) -> FileType:
    if not data:
        raise UploadRejectedError("The file is empty.")
    if len(data) > settings.max_upload_bytes:
        limit_mb = settings.max_upload_bytes / (1024 * 1024)
        raise UploadRejectedError(f"The file is larger than {limit_mb:g} MB.")
    extension = PurePath(filename or "").suffix.lower()
    claimed = _EXTENSIONS.get(extension)
    if claimed is None:
        raise UploadRejectedError("Upload a PDF, DOCX or TXT file.")
    detected = detect_file_type(data)
    if detected is None:
        raise UploadRejectedError("The file content is not a valid PDF, DOCX or plain-text file.")
    if detected != claimed:
        raise UploadRejectedError(
            f"The file's content is {detected.upper()}, but its name ends in {extension}."
        )
    if detected == "docx":
        _check_zip_size(data, settings.max_docx_uncompressed_bytes)
    return detected


def _check_zip_size(data: bytes, limit: int) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        if sum(info.file_size for info in zf.infolist()) > limit:
            raise UploadRejectedError("The document is too large once uncompressed.")


def extract_text(data: bytes, file_type: FileType, settings: Settings) -> ExtractedDocument:
    if file_type == "pdf":
        raw, pages = _extract_pdf(data, settings.max_resume_pages)
    elif file_type == "docx":
        raw, pages = _extract_docx(data, settings.max_docx_uncompressed_bytes), None
    else:
        raw, pages = _decode_text(data), None
    text = normalize_text(raw)
    if len(text) < 50:
        hint = " It may be a scanned image; upload a text-based PDF." if file_type == "pdf" else ""
        raise ExtractionError(f"Could not find enough text in this file.{hint}")
    if len(text) > settings.max_resume_chars:
        raise ExtractionError("This document is too long to be a resume.")
    return ExtractedDocument(text=text, page_count=pages)


def _extract_pdf(data: bytes, max_pages: int) -> tuple[str, int]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("The PDF is password-protected. Upload an unprotected copy.")
        page_count = len(reader.pages)
        if page_count > max_pages:
            raise ExtractionError(f"The PDF has {page_count} pages; the limit is {max_pages}.")
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except ExtractionError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError, OSError) as exc:
        raise ExtractionError("The PDF could not be read. It may be damaged.") from exc
    return text, page_count


def _extract_docx(data: bytes, max_uncompressed: int) -> str:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    _check_zip_size(data, max_uncompressed)
    try:
        document = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, KeyError, ValueError, zipfile.BadZipFile) as exc:
        raise ExtractionError("The DOCX file could not be read. It may be damaged.") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                if cell.text not in cells:  # merged cells repeat their text
                    cells.append(cell.text)
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def _decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ExtractionError("The text file is not valid UTF-8.") from exc
