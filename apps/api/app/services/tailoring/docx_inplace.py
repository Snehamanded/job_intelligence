"""Tailored resume as DOCX, edited inside the user's own DOCX so its format is kept.

Only what tailoring changed is touched: reworded bullets (keeping the paragraph's formatting and
its bold or italic phrases), bullet order, the summary and the skill order. Returns None when a
change can't be placed, so the caller falls back to the template.
"""

import copy
import io
import re
from typing import Any

import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from app.schemas.tailoring import ResumeDocument
from app.services.tailoring.pdf_render import norm, reorder_skills

_SUMMARY = r"summary|profile|objective|about"
_SKILLS = r"skill|technolog|tools|competenc|expertise"


def _is_heading(p: Paragraph) -> bool:
    text = p.text.strip()
    if not text or len(text) > 40 or not re.fullmatch(r"[A-Za-z][A-Za-z &/,'-]+", text):
        return False
    style = (p.style.name if p.style is not None else "") or ""
    runs = [r for r in p.runs if r.text.strip()]
    return (
        style.lower().startswith(("heading", "title")) or text.isupper()
        or (bool(runs) and all(r.bold for r in runs))
    )  # fmt: skip


def _section(paras: list[Paragraph], pattern: str) -> list[int] | None:
    for i, p in enumerate(paras):
        if _is_heading(p) and re.search(pattern, p.text, re.I):
            out = []
            for j in range(i + 1, len(paras)):
                if _is_heading(paras[j]):
                    break
                if paras[j].text.strip():
                    out.append(j)
            return out
    return None


def _find(paras: list[Paragraph], text: str, used: set[int]) -> int | None:
    target = norm(text)
    for i, p in enumerate(paras):
        if i not in used and target and norm(p.text) == target:
            return i
    return None


def _rewrite(p: Paragraph, text: str) -> None:
    """Replace the paragraph's text, keeping its run formatting and differently styled phrases."""
    runs = [r for r in p.runs if r.text]
    if not runs:
        p.add_run(text)
        return

    def key(r: object) -> tuple[bool, bool]:
        return (bool(getattr(r, "bold", False)), bool(getattr(r, "italic", False)))

    counts: dict[tuple[bool, bool], int] = {}
    for r in runs:
        counts[key(r)] = counts.get(key(r), 0) + len(r.text)
    plain = {k: n for k, n in counts.items() if not k[0]} or counts
    base_key = max(plain, key=lambda k: plain[k])
    props: dict[tuple[bool, bool], Any] = {}
    for r in runs:
        props.setdefault(key(r), r._r.rPr)
    phrases = [(r.text.strip(" ,.;:"), key(r)) for r in runs if key(r) != base_key]
    styles = [base_key] * len(text)
    for phrase, k in sorted((x for x in phrases if len(x[0]) >= 2), key=lambda x: -len(x[0])):
        for m in re.finditer(r"(?<!\w)" + re.escape(phrase) + r"(?!\w)", text):
            if all(s == base_key for s in styles[m.start() : m.end()]):
                styles[m.start() : m.end()] = [k] * (m.end() - m.start())
    for r in list(p._p.iter(qn("w:r"))):
        r.getparent().remove(r)
    start = 0
    for i in range(1, len(text) + 1):
        if i == len(text) or styles[i] != styles[start]:
            run = p.add_run(text[start:i])
            rpr = props.get(styles[start])
            if rpr is not None:
                run._r.insert(0, copy.deepcopy(rpr))
            start = i


def tailor_docx(original: bytes, base: ResumeDocument, doc: ResumeDocument) -> bytes | None:
    try:
        document = docx.Document(io.BytesIO(original))
    except Exception:
        return None
    paras = [Paragraph(p, document._body) for p in document.element.body.iter(qn("w:p"))]
    used: set[int] = set()
    rewrites: dict[int, str] = {}
    for base_role in base.experience:
        role = next((r for r in doc.experience if r.id == base_role.id), None)
        if role is None:
            continue
        if [b.id for b in role.bullets] == [b.id for b in base_role.bullets] and not any(
            b.ai_changed for b in role.bullets
        ):
            continue
        found: dict[str, int] = {}
        for bullet in base_role.bullets:
            index = _find(paras, bullet.original_text, used)
            if index is None:
                return None
            found[bullet.id] = index
            used.add(index)
        slots = sorted(found.values())
        sources = {i: copy.deepcopy(paras[i]._p) for i in slots}
        for n, slot in enumerate(slots):
            element = paras[slot]._p
            if n < len(role.bullets):
                bullet = role.bullets[n]
                moved = copy.deepcopy(sources[found[bullet.id]])
                element.addprevious(moved)
                paras[slot] = Paragraph(moved, paras[slot]._parent)
                if bullet.ai_changed:
                    rewrites[slot] = bullet.text
            element.getparent().remove(element)  # bullets dropped from the role go with it
    if doc.summary:
        section = _section(paras, _SUMMARY)
        if not section:
            return None
        rewrites[section[0]] = " ".join(s.text for s in doc.summary)
        for j in section[1:]:
            paras[j]._p.getparent().remove(paras[j]._p)
    if [s.name for s in doc.skills] != [s.name for s in base.skills]:
        section = _section(paras, _SKILLS)
        if section is None:
            return None
        for j in section:
            text = reorder_skills(paras[j].text, base, doc)
            if text is not None:
                rewrites[j] = text
    for index, text in rewrites.items():
        _rewrite(paras[index], text)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()
