"""ATS-friendly DOCX: one column, standard fonts and headings, no tables or text boxes."""

import io

import docx
from docx.shared import Pt

from app.schemas.tailoring import DocEducation, ResumeDocument


def education_line(e: DocEducation) -> str:
    # Skip the field when the degree already names it ("B.E. in Computer Science").
    field = e.field_of_study
    if field and e.degree and field.lower() in e.degree.lower():
        field = None
    return ", ".join(x for x in [e.institution, e.degree, field, e.date_text] if x)


def render_docx(doc: ResumeDocument) -> bytes:
    out = docx.Document()
    style = out.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    c = doc.contact
    if c.name:
        out.add_heading(c.name, level=0)
    contact_line = " | ".join(x for x in [c.email, c.phone, c.location, *c.links] if x)
    if contact_line:
        out.add_paragraph(contact_line)

    if doc.summary:
        out.add_heading("Summary", level=1)
        out.add_paragraph(" ".join(s.text for s in doc.summary))
    if doc.skills:
        out.add_heading("Skills", level=1)
        out.add_paragraph(", ".join(s.name for s in doc.skills))
    if doc.experience:
        out.add_heading("Experience", level=1)
        for role in doc.experience:
            p = out.add_paragraph()
            p.add_run(f"{role.title}, {role.company}").bold = True
            details = " | ".join(x for x in [role.location, role.date_text] if x)
            if details:
                p.add_run(f"  {details}")
            for bullet in role.bullets:
                out.add_paragraph(bullet.text, style="List Bullet")
    if doc.projects:
        out.add_heading("Projects", level=1)
        for project in doc.projects:
            p = out.add_paragraph()
            p.add_run(project.name).bold = True
            if project.description:
                p.add_run(f": {project.description}")
            if project.technologies:
                out.add_paragraph("Technologies: " + ", ".join(project.technologies))
    if doc.education:
        out.add_heading("Education", level=1)
        for e in doc.education:
            out.add_paragraph(education_line(e))
    if doc.certifications:
        out.add_heading("Certifications", level=1)
        for cert in doc.certifications:
            out.add_paragraph(
                cert.name + (f" ({cert.issuer})" if cert.issuer else ""), style="List Bullet"
            )

    buffer = io.BytesIO()
    out.save(buffer)
    return buffer.getvalue()
