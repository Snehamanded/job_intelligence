import io
import json
from typing import Any

import docx
from fastapi.testclient import TestClient

from app.ai.providers.base import AIProviderError
from app.ai.providers.fake import FakeAIProvider
from app.schemas.tailoring import Change, DocBullet, DocRole, DocSkill, ResumeDocument
from app.services.tailoring.checks import check_rewrite, check_summary_sentence
from app.services.tailoring.document import apply
from app.services.tailoring.rules import rule_changes
from tests.conftest import InlineQueue, give_consent, register, upload
from tests.helpers import fixture_text

ROLE = "Software Engineer\nAcme\nBuilt REST APIs in FastAPI serving 2,000 daily users\nReduced report time by 40% with PostgreSQL indexes"
JOB = {
    "title": "Backend Engineer",
    "company": "Acme Payments",
    "location": "Bengaluru, India",
    "description": "We need PostgreSQL, Python and FastAPI to build REST APIs. Docker, Kubernetes and Go are a plus. 1-3 years of experience. Full-time.",
}

# --- code checks ---------------------------------------------------------------------------


def test_check_rewrite_rules() -> None:
    original = "Built REST APIs in FastAPI serving 2,000 daily users"
    assert (
        check_rewrite(
            original, "Designed and built RESTful APIs with FastAPI for 2000 daily users", ROLE
        )
        == []
    )
    assert any(
        "60" in i
        for i in check_rewrite(original, "Built FastAPI APIs serving 60% more users", ROLE)
    )
    assert any(
        "Kubernetes" in i
        for i in check_rewrite(original, "Built REST APIs in FastAPI on Kubernetes", ROLE)
    )
    # A technology from elsewhere in the same role is fine.
    assert check_rewrite(original, "Built REST APIs in FastAPI backed by PostgreSQL", ROLE) == []
    assert any(
        "longer" in i
        for i in check_rewrite("Wrote tests", "Wrote tests " + "carefully " * 20, ROLE)
    )
    assert check_rewrite(original, "   ", ROLE) == ["The rewrite is empty"]


def test_check_summary_sentence() -> None:
    texts = {"s0": "Python", "r0b0": "Built REST APIs in FastAPI serving 2,000 daily users"}
    assert (
        check_summary_sentence(
            "Python developer building REST APIs in FastAPI.", ["s0", "r0b0"], texts
        )
        == []
    )
    assert check_summary_sentence("Python developer.", [], texts) == [
        "Doesn't cite any item on your resume"
    ]
    assert check_summary_sentence("Python developer.", ["nope"], texts) == [
        "Doesn't cite any item on your resume"
    ]
    assert any("5" in i for i in check_summary_sentence("5 years of Python.", ["s0"], texts))
    assert any(
        "AWS" in i for i in check_summary_sentence("Python developer on AWS.", ["s0"], texts)
    )


def _doc() -> ResumeDocument:
    return ResumeDocument(
        skills=[DocSkill(id="s0", name="React"), DocSkill(id="s1", name="Python"), DocSkill(id="s2", name="PostgreSQL")],
        experience=[
            DocRole(
                id="r0", title="Engineer", company="Acme",
                bullets=[
                    DocBullet(id="r0b0", text="Built dashboards in React", original_text="Built dashboards in React", source_span=None),
                    DocBullet(id="r0b1", text="Tuned PostgreSQL queries", original_text="Tuned PostgreSQL queries", source_span=None),
                ],
            )
        ],
    )  # fmt: skip


def test_rule_changes_reorder_by_job_skills() -> None:
    changes = {c.id: c for c in rule_changes(_doc(), ["PostgreSQL", "Python"])}
    assert changes["rules-skills"].skill_ids == ["s2", "s1", "s0"]
    assert changes["rules-bullets-r0"].bullet_ids == ["r0b1", "r0b0"]
    assert rule_changes(_doc(), []) == []


def test_apply_never_applies_unsupported_changes() -> None:
    bad = Change(id="x", kind="bullet_rewrite", source="ai", status="unsupported", decision="accepted",
                 title="t", bullet_id="r0b0", after="Led 10 people")  # fmt: skip
    good = Change(id="y", kind="bullet_rewrite", source="ai", status="proposed", decision="accepted",
                  title="t", bullet_id="r0b1", after="Optimised PostgreSQL queries")  # fmt: skip
    doc = apply(_doc(), [bad, good])
    texts = [b.text for b in doc.experience[0].bullets]
    assert texts == ["Built dashboards in React", "Optimised PostgreSQL queries"]
    assert (
        doc.experience[0].bullets[1].ai_changed
        and doc.experience[0].bullets[1].original_text == "Tuned PostgreSQL queries"
    )


# --- end to end with a fake LLM that invents claims ------------------------------------------

SUGGESTIONS = {
    "skills_order": ["s3", "s0", "s6", "s7", "bogus"],
    "roles": [
        {"role_id": "r0", "bullet_ids": ["r0b1", "r0b0", "r0b2", "nope"]},
        {"role_id": "r9", "bullet_ids": []},
    ],
    "rewrites": [
        {
            "bullet_id": "r0b0",
            "text": "Designed and built REST APIs with FastAPI serving 2,000 daily users",
            "reason": "Stronger verb",
        },
        {
            "bullet_id": "r0b1",
            "text": "Cut report generation time by 60% with PostgreSQL indexing",
            "reason": "Punchier",
        },
        {
            "bullet_id": "r1b0",
            "text": "Developed internal dashboards in React deployed on Kubernetes",
            "reason": "Match the job",
        },
        {
            "bullet_id": "r0b2",
            "text": "Led a team writing unit tests with pytest and GitHub Actions CI",
            "reason": "Leadership",
        },
        {"bullet_id": "zz", "text": "Invented bullet", "reason": ""},
    ],
    "summary": [
        {
            "text": "Python backend developer building REST APIs with FastAPI.",
            "sources": ["s0", "s3", "r0b0"],
        }
    ],
}
VERIFY = {
    "results": [
        {"change_id": "ai-rewrite-r0b0", "supported": True, "problems": []},
        {
            "change_id": "ai-rewrite-r0b2",
            "supported": False,
            "problems": ["Leading a team is not stated"],
        },
        {"change_id": "ai-summary", "supported": True, "problems": []},
    ]
}


def setup(client: TestClient, queue: InlineQueue, *, consent: bool) -> tuple[dict[str, str], str]:
    headers = register(client)
    upload(
        client, headers, "sneha.txt", fixture_text("sneha_backend.txt").encode()
    )  # basic parsing
    job_id: str = client.post("/api/jobs/import", json=JOB, headers=headers).json()["id"]
    if consent:
        queue.paused = True  # don't trigger AI scoring; this test is about tailoring
        give_consent(client, headers)
        queue.paused = False
    return headers, job_id


def start(client: TestClient, headers: dict[str, str], job_id: str) -> dict[str, Any]:
    resp = client.post("/api/tailoring", json={"job_id": job_id}, headers=headers)
    assert resp.status_code == 202, resp.text
    detail: dict[str, Any] = client.get(f"/api/tailoring/{resp.json()['id']}").json()
    return detail


def test_ai_tailoring_end_to_end(
    client: TestClient, queue: InlineQueue, fake_ai: FakeAIProvider
) -> None:
    headers, job_id = setup(client, queue, consent=True)
    fake_ai.responses.extend([json.dumps(SUGGESTIONS), json.dumps(VERIFY)])
    v = start(client, headers, job_id)
    assert (v["status"], v["method"]) == ("ready", "ai")

    changes = {c["id"]: c for c in v["changes"]}
    assert changes["ai-rewrite-r0b0"]["status"] == "proposed"
    assert changes["ai-rewrite-r0b1"]["status"] == "unsupported"
    assert "60" in changes["ai-rewrite-r0b1"]["issues"][0]
    assert changes["ai-rewrite-r1b0"]["status"] == "unsupported"
    assert "Kubernetes" in changes["ai-rewrite-r1b0"]["issues"][0]
    assert changes["ai-rewrite-r0b2"]["issues"] == ["Second check: Leading a team is not stated"]
    assert changes["ai-summary"]["status"] == "proposed"
    assert changes["ai-skills"]["skill_ids"] == ["s3", "s0", "s6", "s7"]  # bogus id dropped
    assert changes["ai-bullets-r0"]["bullet_ids"] == ["r0b1", "r0b0", "r0b2"]
    assert "ai-rewrite-zz" not in changes
    assert {c["source"] for c in v["changes"]} == {"ai", "rules"}

    # The second pass only saw wording that passed the code checks; job text is delimited data.
    suggest_call, verify_call = fake_ai.calls[-2:]
    assert "<job_posting>" in suggest_call.prompt
    assert "ai-rewrite-r0b0" in verify_call.prompt and "ai-rewrite-r0b1" not in verify_call.prompt

    url = f"/api/tailoring/{v['id']}/decisions"
    refused = client.patch(
        url, json={"decisions": {"ai-rewrite-r0b1": "accepted"}}, headers=headers
    )
    assert refused.status_code == 422 and "can't be accepted" in refused.json()["detail"]
    ok = client.patch(url, json={"decisions": {"ai-rewrite-r0b0": "accepted", "ai-summary": "accepted",
                                                "ai-skills": "accepted", "rules-skills": "accepted"}},
                      headers=headers).json()  # fmt: skip
    decided = {c["id"]: c["decision"] for c in ok["changes"]}
    assert (
        decided["rules-skills"] == "accepted" and decided["ai-skills"] == "rejected"
    )  # one per target
    assert ok["preview"]["experience"][0]["bullets"][0]["text"].startswith("Designed and built")

    saved = client.post(
        f"/api/tailoring/{v['id']}/save", json={"name": "Acme backend"}, headers=headers
    ).json()
    assert (saved["status"], saved["name"]) == ("saved", "Acme backend")
    doc = saved["preview"]
    bullet = doc["experience"][0]["bullets"][0]
    assert (
        bullet["ai_changed"]
        and bullet["original_text"] == "Built REST APIs in FastAPI serving 2,000 daily users"
    )
    resume_text = client.get("/api/me/export").json()["resumes"][0]["extracted_text"]
    span = bullet["source_span"]
    assert (
        resume_text[span["start"] : span["end"]] == bullet["original_text"]
    )  # traces to the resume
    all_text = json.dumps(doc)
    assert "60%" not in all_text and "Kubernetes" not in all_text and "Led a team" not in all_text
    profile_skills = {s["name"] for s in client.get("/api/profile").json()["data"]["skills"]}
    assert {s["name"] for s in doc["skills"]} <= profile_skills  # never related or missing skills
    assert doc["summary"][0]["sources"] == ["s0", "s3", "r0b0"]
    assert any(
        g["name"] == "Kubernetes" for g in saved["skill_gaps"]
    )  # gaps shown beside, not in, resume

    # Saved versions are immutable.
    assert (
        client.patch(
            url, json={"decisions": {"ai-summary": "rejected"}}, headers=headers
        ).status_code
        == 409
    )
    assert (
        client.post(f"/api/tailoring/{v['id']}/save", json={}, headers=headers).status_code == 409
    )

    resp = client.get(f"/api/tailoring/{v['id']}/docx")
    assert (
        resp.status_code == 200
        and 'filename="Acme-backend.docx"' in resp.headers["content-disposition"]
    )
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(resp.content)).paragraphs)
    assert "Designed and built REST APIs" in text and "60%" not in text and "Sneha Rao" in text

    exported = client.get("/api/me/export").json()["resume_versions"]
    assert exported[0]["name"] == "Acme backend" and exported[0]["content"]["summary"]


def test_verifier_unavailable_blocks_rewrites(
    client: TestClient, queue: InlineQueue, fake_ai: FakeAIProvider
) -> None:
    headers, job_id = setup(client, queue, consent=True)
    fake_ai.responses.extend([json.dumps(SUGGESTIONS), AIProviderError("down", retryable=False)])
    v = start(client, headers, job_id)
    rewrite = next(c for c in v["changes"] if c["id"] == "ai-rewrite-r0b0")
    assert rewrite["status"] == "unsupported" and "Couldn't be checked" in rewrite["issues"][0]
    reorder = next(c for c in v["changes"] if c["id"] == "ai-skills")
    assert reorder["status"] == "proposed"  # reordering adds no claims


def test_rule_based_without_consent(
    client: TestClient, queue: InlineQueue, fake_ai: FakeAIProvider
) -> None:
    headers, job_id = setup(client, queue, consent=False)
    v = start(client, headers, job_id)
    assert (v["status"], v["method"]) == ("ready", "rules")
    assert "only rule-based" in v["notice"]
    assert fake_ai.calls == []
    assert {c["kind"] for c in v["changes"]} <= {"skills_order", "bullets_order"}
    skills = next(c for c in v["changes"] if c["id"] == "rules-skills")
    client.patch(
        f"/api/tailoring/{v['id']}/decisions",
        json={"decisions": {skills["id"]: "accepted"}},
        headers=headers,
    )
    saved = client.post(f"/api/tailoring/{v['id']}/save", json={}, headers=headers).json()
    assert [s["name"] for s in saved["preview"]["skills"]][:3] == [
        "PostgreSQL",
        "Python",
        "FastAPI",
    ]
    assert client.get(f"/api/tailoring?job_id={job_id}").json()[0]["status"] == "saved"


def test_requires_resume_and_own_job(client: TestClient) -> None:
    headers = register(client, email="a@example.com")
    job_id = client.post("/api/jobs/import", json=JOB, headers=headers).json()["id"]
    no_resume = client.post("/api/tailoring", json={"job_id": job_id}, headers=headers)
    assert no_resume.status_code == 422 and "Upload your resume" in no_resume.json()["detail"]
    client.post("/api/auth/logout", headers=headers)
    other = register(client, email="b@example.com")
    assert client.post("/api/tailoring", json={"job_id": job_id}, headers=other).status_code == 404


def test_versions_are_private(client: TestClient, queue: InlineQueue) -> None:
    headers, job_id = setup(client, queue, consent=False)
    v = start(client, headers, job_id)
    client.post("/api/auth/logout", headers=headers)
    other = register(client, email="other@example.com")
    assert client.get(f"/api/tailoring/{v['id']}").status_code == 404
    assert client.get(f"/api/tailoring/{v['id']}/docx").status_code == 404
    assert client.delete(f"/api/tailoring/{v['id']}", headers=other).status_code == 404
    assert client.get("/api/tailoring").json() == []


def test_education_line_skips_repeated_field() -> None:
    from app.schemas.tailoring import DocEducation
    from app.services.tailoring.docx_render import education_line

    e = DocEducation(
        institution="RV College",
        degree="B.E. in Computer Science",
        field_of_study="Computer Science",
        date_text="2021 - 2025",
    )
    assert education_line(e) == "RV College, B.E. in Computer Science, 2021 - 2025"
    e2 = DocEducation(institution="RV College", degree="B.E.", field_of_study="Computer Science")
    assert education_line(e2) == "RV College, B.E., Computer Science"
