"""Tailored resume downloads keep the uploaded resume's format. All content is fictional."""

import io
import json
import logging
from typing import Any

import docx
import pdfplumber
import pytest
from fastapi.testclient import TestClient

from app.ai.providers.fake import FakeAIProvider
from app.schemas.tailoring import DocBullet, DocRole, DocSkill, ResumeDocument
from app.services.resume.layout import extract_layout, layout_text
from app.services.tailoring.pdf_render import render_pdf, reorder_skills
from tests.conftest import InlineQueue, give_consent, register, upload
from tests.helpers import fixture_text, make_styled_pdf
from tests.test_tailoring import JOB, SUGGESTIONS, VERIFY

BOLD = ("REST APIs", "40%")
ACCEPT = {"ai-rewrite-r0b0": "accepted", "ai-summary": "accepted", "rules-skills": "accepted",
          "ai-bullets-r0": "accepted"}  # fmt: skip


def _tailor(
    client: TestClient, queue: InlineQueue, fake_ai: FakeAIProvider, name: str, data: bytes
) -> tuple[dict[str, str], str]:
    headers = register(client)
    upload(client, headers, name, data)
    job_id = client.post("/api/jobs/import", json=JOB, headers=headers).json()["id"]
    queue.paused = True
    give_consent(client, headers)
    queue.paused = False
    fake_ai.responses.extend([json.dumps(SUGGESTIONS), json.dumps(VERIFY)])
    version_id: str = client.post(
        "/api/tailoring", json={"job_id": job_id}, headers=headers
    ).json()["id"]
    client.patch(f"/api/tailoring/{version_id}/decisions", json={"decisions": ACCEPT},
                 headers=headers)  # fmt: skip
    saved = client.post(f"/api/tailoring/{version_id}/save", json={"name": "Acme backend"},
                        headers=headers)  # fmt: skip
    assert saved.json()["status"] == "saved", saved.text
    return headers, version_id


def _words(data: bytes) -> list[dict[str, Any]]:
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return [w for p in pdf.pages for w in p.extract_words(extra_attrs=["fontname"])]


def test_pdf_download_keeps_the_uploaded_pdf_format(
    client: TestClient,
    queue: InlineQueue,
    fake_ai: FakeAIProvider,
    caplog: pytest.LogCaptureFixture,
) -> None:
    original = make_styled_pdf(fixture_text("sneha_backend.txt"), bold=BOLD)
    _, version_id = _tailor(client, queue, fake_ai, "resume.pdf", original)
    with caplog.at_level(logging.INFO):
        resp = client.get(f"/api/tailoring/{version_id}/pdf")
    assert resp.status_code == 200 and resp.headers["content-type"] == "application/pdf"
    assert [
        getattr(r, "kept_format", None) for r in caplog.records if r.msg == "resume_pdf_rendered"
    ] == [True]
    assert 'filename="Acme-backend.pdf"' in resp.headers["content-disposition"]

    layout = extract_layout(resp.content)
    assert layout is not None
    text = layout_text(layout)
    # Accepted changes: new order, reworded bullet, summary.
    order = [text.index(s) for s in ("Reduced report generation", "Designed and built REST APIs",
                                     "Wrote unit tests")]  # fmt: skip
    assert order == sorted(order)
    assert "Built REST APIs in FastAPI" not in text
    assert "Python backend developer building REST APIs with FastAPI." in text
    # Everything else is kept, in the same style: same rules, small-caps name, bold phrases.
    before = extract_layout(original)
    assert before is not None and len(layout.rules) == len(before.rules) == 6
    assert "AWS Certified Cloud Practitioner" in text and "Nimbus Labs" in text
    words = _words(resp.content)
    assert any(w["text"] == "Rao" and "Caps" in w["fontname"] for w in words)
    designed = next(w for w in words if w["text"] == "Designed")
    line = [w for w in words if abs(w["top"] - designed["top"]) < 1]
    assert {w["text"]: "Bold" in w["fontname"] for w in line if w["text"] in ("Designed", "REST")
            } == {"Designed": False, "REST": True}  # fmt: skip
    # Built on the server: no browser header or footer (date, URL, page title).
    assert "http" not in text and "localhost" not in text


def test_pdf_download_uses_the_template_without_a_pdf_layout(
    client: TestClient, queue: InlineQueue, fake_ai: FakeAIProvider
) -> None:
    text_resume = fixture_text("sneha_backend.txt").encode()
    _, version_id = _tailor(client, queue, fake_ai, "resume.txt", text_resume)
    resp = client.get(f"/api/tailoring/{version_id}/pdf")
    assert resp.status_code == 200
    layout = extract_layout(resp.content)
    assert layout is not None
    text = layout_text(layout)
    assert "Sneha Rao" in text and "Designed and built REST APIs" in text
    assert "Experience" in text and "Certifications" in text


def test_docx_download_edits_the_uploaded_docx(
    client: TestClient, queue: InlineQueue, fake_ai: FakeAIProvider
) -> None:
    document = docx.Document()
    for line in fixture_text("sneha_backend.txt").splitlines():
        if line in ("Summary", "Skills", "Experience", "Projects", "Education", "Certifications"):
            document.add_heading(line, level=1)
        elif line.startswith("• "):
            p = document.add_paragraph(style="List Bullet")
            for part in line[2:].partition("REST APIs"):
                if part:
                    p.add_run(part).bold = part == "REST APIs"
        else:
            document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    _, version_id = _tailor(client, queue, fake_ai, "resume.docx", buffer.getvalue())

    resp = client.get(f"/api/tailoring/{version_id}/docx")
    assert resp.status_code == 200
    out = docx.Document(io.BytesIO(resp.content))
    paras = [p for p in out.paragraphs if p.text.strip()]
    texts = [p.text for p in paras]
    designed = next(p for p in paras if p.text.startswith("Designed and built REST APIs"))
    assert designed.style is not None and designed.style.name == "List Bullet"  # format kept
    assert [r.text for r in designed.runs if r.bold] == ["REST APIs"]
    i = texts.index
    assert (
        i("Reduced report generation time by 40% by adding PostgreSQL indexes")
        < i(designed.text)
        < i("Wrote unit tests with pytest and set up GitHub Actions CI")
    )
    assert "Python backend developer building REST APIs with FastAPI." in texts
    assert out.paragraphs[0].text == "Sneha Rao" and "Certifications" in texts


def _doc(bullets: list[str]) -> ResumeDocument:
    return ResumeDocument(
        skills=[DocSkill(id=f"s{i}", name=n) for i, n in enumerate(["Python", "SQL", "Docker"])],
        experience=[
            DocRole(id="r0", title="Engineer", company="Acme",
                    bullets=[DocBullet(id=f"r0b{i}", text=t, original_text=t, source_span=None)
                             for i, t in enumerate(bullets)])
        ],
    )  # fmt: skip


def test_unplaceable_change_falls_back_to_the_template() -> None:
    original = make_styled_pdf(fixture_text("sneha_backend.txt"), bold=BOLD)
    layout = extract_layout(original)
    base = _doc(["This bullet is not in the PDF at all"])
    doc = base.model_copy(deep=True)
    doc.experience[0].bullets[0].text, doc.experience[0].bullets[0].ai_changed = "New", True
    _, kept = render_pdf(layout, base, doc)
    assert kept is False
    _, kept = render_pdf(layout, base, base)  # nothing to place: the original is kept
    assert kept is True


def test_reorder_skills_keeps_labels_and_unranked_items() -> None:
    base = _doc([])
    doc = base.model_copy(deep=True)
    doc.skills = [doc.skills[2], doc.skills[0]]  # Docker first; SQL dropped
    assert reorder_skills("Languages: Python, SQL, Go", base, doc) == "Languages: Python, Go"
    assert reorder_skills("Tools: Git, Python, Docker", base, doc) == "Tools: Git, Docker, Python"
    assert reorder_skills("Tools: Git, Make", base, doc) is None
